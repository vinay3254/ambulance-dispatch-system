"""
Create Hackathon Demo Documentation for Ambulance Dispatch System
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from datetime import datetime

doc = Document()

# ============= STYLES =============
title_style = doc.styles['Title']
title_style.font.size = Pt(28)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(220, 53, 69)

heading1_style = doc.styles['Heading 1']
heading1_style.font.size = Pt(18)
heading1_style.font.color.rgb = RGBColor(0, 123, 255)

heading2_style = doc.styles['Heading 2']
heading2_style.font.size = Pt(14)
heading2_style.font.color.rgb = RGBColor(40, 167, 69)

# ============= TITLE PAGE =============
title = doc.add_paragraph()
title_run = title.add_run("🚑 AI-POWERED INTELLIGENT AMBULANCE DISPATCH SYSTEM")
title_run.bold = True
title_run.font.size = Pt(24)
title_run.font.color.rgb = RGBColor(220, 53, 69)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph()
sub_run = subtitle.add_run("Priority Routing & Green Corridor with Auto Hospital Allocation")
sub_run.font.size = Pt(16)
sub_run.font.color.rgb = RGBColor(108, 117, 125)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()
doc.add_paragraph()

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info.add_run("HACKATHON DEMO GUIDE & VIDEO SCRIPT\n\n").bold = True
info.add_run(f"Document Generated: {datetime.now().strftime('%B %d, %Y')}\n")
info.add_run("Version: 1.0\n\n")
info.add_run("━" * 50)

doc.add_page_break()

# ============= TABLE OF CONTENTS =============
doc.add_heading("📋 TABLE OF CONTENTS", level=1)
toc_items = [
    "1. Project Overview & Problem Statement",
    "2. System Architecture",
    "3. Quick Start Guide",
    "4. Demo Workflow (Video Script)",
    "5. Feature Highlights",
    "6. API Endpoints Reference",
    "7. Technology Stack",
    "8. Future Enhancements",
    "9. Team Information"
]
for item in toc_items:
    p = doc.add_paragraph(item)
    p.paragraph_format.left_indent = Inches(0.5)

doc.add_page_break()

# ============= SECTION 1: OVERVIEW =============
doc.add_heading("1. PROJECT OVERVIEW & PROBLEM STATEMENT", level=1)

doc.add_heading("The Golden Hour Crisis", level=2)
doc.add_paragraph(
    "Every year in India, thousands of preventable deaths occur not because of lack of medical "
    "facilities — but because the right patient never reaches the right hospital in time. "
    "The golden hour — the critical 60-minute window after trauma, cardiac arrest, or stroke "
    "where emergency intervention saves lives — is routinely lost."
)

doc.add_heading("Key Problems We Solve", level=2)
problems = [
    ("No unified dispatch protocol", "108, 102, 1099 operate in silos with zero coordination"),
    ("Manual dispatcher decisions", "Based on guesswork, not real-time GPS/traffic data"),
    ("Nearest ambulance myth", "Dispatchers use outdated zone maps, not live positions"),
    ("Hospital capacity blindness", "Ambulances rush to full ICU/trauma facilities"),
    ("Traffic signal unawareness", "No smart green corridor activation in most cities"),
    ("Rural-urban gap", "Tier 2/3 cities have near-zero structured EMS coverage"),
]
for title, desc in problems:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(f"{title}: ").bold = True
    p.add_run(desc)

doc.add_heading("Our Solution", level=2)
doc.add_paragraph(
    "An end-to-end AI-driven ambulance dispatch platform featuring real-time optimal dispatch, "
    "dynamic priority routing with automated green corridor activation, intelligent hospital "
    "matching based on live capacity, and predictive fleet positioning."
)

doc.add_page_break()

# ============= SECTION 2: ARCHITECTURE =============
doc.add_heading("2. SYSTEM ARCHITECTURE", level=1)

doc.add_heading("Microservices Components", level=2)
components = [
    ("API Gateway", "Central routing hub for all services"),
    ("Dispatch Engine", "Real-time ambulance selection algorithm"),
    ("Routing Service", "OSRM-based navigation with traffic integration"),
    ("Hospital Matcher", "Multi-factor scoring for optimal facility selection"),
    ("Siren Detection", "Audio FFT + Camera OCR for ambulance detection"),
    ("Traffic Controller", "Green corridor activation at intersections"),
    ("WebSocket Server", "Real-time updates to all connected clients"),
    ("Prediction Engine", "ML-based demand forecasting for fleet positioning"),
]
for name, desc in components:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(f"{name}: ").bold = True
    p.add_run(desc)

doc.add_heading("Data Flow", level=2)
doc.add_paragraph(
    "Emergency Call → Dispatch Engine → Ambulance Assignment → Route Calculation → "
    "Green Corridor Activation → Hospital Matching → Real-time Tracking → Arrival Confirmation"
)

doc.add_page_break()

# ============= SECTION 3: QUICK START =============
doc.add_heading("3. QUICK START GUIDE", level=1)

doc.add_heading("Prerequisites", level=2)
prereqs = ["Node.js v18+", "Python 3.11+", "Modern web browser (Chrome/Edge recommended)"]
for p in prereqs:
    doc.add_paragraph(p, style='List Bullet')

doc.add_heading("Starting the System", level=2)
doc.add_paragraph("Step 1: Open Terminal/PowerShell in the project directory")
doc.add_paragraph("Step 2: Start the backend server:")
code1 = doc.add_paragraph()
code1.add_run("cd ambulance-dispatch-system").font.name = 'Consolas'
doc.add_paragraph()
code2 = doc.add_paragraph()
code2.add_run("node demo-server.js").font.name = 'Consolas'

doc.add_paragraph("\nStep 3: In a new terminal, start the frontend:")
code3 = doc.add_paragraph()
code3.add_run("cd frontend/dispatcher-dashboard").font.name = 'Consolas'
doc.add_paragraph()
code4 = doc.add_paragraph()
code4.add_run("npm run dev").font.name = 'Consolas'

doc.add_heading("Access Points", level=2)
table = doc.add_table(rows=4, cols=2)
table.style = 'Table Grid'
cells = table.rows[0].cells
cells[0].text = "Component"
cells[1].text = "URL"
cells[0].paragraphs[0].runs[0].bold = True
cells[1].paragraphs[0].runs[0].bold = True

data = [
    ("Dispatcher Dashboard", "http://localhost:5173"),
    ("API Server", "http://localhost:3000"),
    ("Login Credentials", "dispatcher@demo.com / password123"),
]
for i, (comp, url) in enumerate(data, 1):
    cells = table.rows[i].cells
    cells[0].text = comp
    cells[1].text = url

doc.add_page_break()

# ============= SECTION 4: DEMO WORKFLOW =============
doc.add_heading("4. DEMO WORKFLOW (VIDEO SCRIPT)", level=1)

doc.add_paragraph(
    "Follow this step-by-step script to record your hackathon demo video. "
    "Estimated total duration: 5-7 minutes."
).italic = True

doc.add_heading("🎬 Scene 1: Introduction (30 seconds)", level=2)
doc.add_paragraph("NARRATION:", style='List Bullet').runs[0].bold = True
doc.add_paragraph(
    '"Welcome to our AI-Powered Intelligent Ambulance Dispatch System. Every year, '
    'thousands die because the right patient never reaches the right hospital in time. '
    'Our solution addresses this golden hour crisis with real-time dispatch, smart routing, '
    'and automated green corridor activation."'
)
doc.add_paragraph("ACTION: Show the login page at http://localhost:5173", style='List Bullet')

doc.add_heading("🎬 Scene 2: Login & Dashboard Overview (45 seconds)", level=2)
doc.add_paragraph("ACTION:", style='List Bullet').runs[0].bold = True
steps = [
    "Enter email: dispatcher@demo.com",
    "Enter password: password123", 
    "Click Login button",
    "Pan across the dashboard showing: Live Map, Incident Queue, Statistics Cards"
]
for s in steps:
    doc.add_paragraph(f"  • {s}")

doc.add_paragraph("NARRATION:", style='List Bullet').runs[0].bold = True
doc.add_paragraph(
    '"This is our unified command center. Dispatchers see all active emergencies, '
    'ambulance positions, and hospital statuses in real-time. Notice the color-coded '
    'incident markers — red for critical, orange for urgent, yellow for non-urgent."'
)

doc.add_heading("🎬 Scene 3: Incident Queue & Details (45 seconds)", level=2)
doc.add_paragraph("ACTION:", style='List Bullet').runs[0].bold = True
doc.add_paragraph("  • Click on an incident in the queue (left sidebar)")
doc.add_paragraph("  • Show incident details: location, type, severity, timestamp")
doc.add_paragraph("  • Highlight the 'Assign Ambulance' button")

doc.add_paragraph("NARRATION:", style='List Bullet').runs[0].bold = True
doc.add_paragraph(
    '"Each incident shows complete details including GPS coordinates captured automatically '
    'from the caller\'s phone. Our AI has already calculated the optimal ambulance based on '
    'proximity, traffic conditions, and vehicle type."'
)

doc.add_heading("🎬 Scene 4: Ambulance Assignment Wizard (60 seconds)", level=2)
doc.add_paragraph("ACTION:", style='List Bullet').runs[0].bold = True
doc.add_paragraph("  • Click 'Assign Ambulance' on a pending incident")
doc.add_paragraph("  • Step 1: Show recommended ambulances ranked by AI score")
doc.add_paragraph("  • Step 2: Select ambulance, show ETA calculation")
doc.add_paragraph("  • Step 3: Confirm hospital recommendation")
doc.add_paragraph("  • Click Confirm Assignment")

doc.add_paragraph("NARRATION:", style='List Bullet').runs[0].bold = True
doc.add_paragraph(
    '"Our dispatch engine considers multiple factors: true road distance (not straight-line), '
    'current traffic, ambulance type matching (ALS for cardiac, neonatal for infants), '
    'and crew availability. The AI recommends optimal choices but dispatchers retain override control."'
)

doc.add_heading("🎬 Scene 5: Green Corridor / Siren Detection (60 seconds)", level=2)
doc.add_paragraph("ACTION:", style='List Bullet').runs[0].bold = True
doc.add_paragraph("  • Open new browser tab: http://localhost:3000/api/siren-detection/traffic")
doc.add_paragraph("  • Show traffic light states (JSON response)")
doc.add_paragraph("  • Use Postman/curl to activate green corridor:")

code_block = doc.add_paragraph()
code_block.add_run(
    'POST http://localhost:3000/api/green-corridor/activate\n'
    '{"ambulance_id": "amb-001", "lane": "NORTH"}'
).font.name = 'Consolas'

doc.add_paragraph("  • Refresh traffic endpoint to show NORTH is now GREEN")

doc.add_paragraph("NARRATION:", style='List Bullet').runs[0].bold = True
doc.add_paragraph(
    '"This is our integrated siren detection and green corridor system. Using audio FFT analysis '
    'and camera-based ambulance detection, intersections automatically detect approaching ambulances '
    'and switch traffic signals to green. Watch as the NORTH lane turns green while others turn red, '
    'creating a clear corridor for the ambulance."'
)

doc.add_heading("🎬 Scene 6: Hospital Matching (45 seconds)", level=2)
doc.add_paragraph("ACTION:", style='List Bullet').runs[0].bold = True
doc.add_paragraph("  • Open: http://localhost:3000/api/hospitals")
doc.add_paragraph("  • Show hospital data with bed availability, specialties")
doc.add_paragraph("  • Highlight ICU beds, trauma bays, specialist on-duty")

doc.add_paragraph("NARRATION:", style='List Bullet').runs[0].bold = True
doc.add_paragraph(
    '"Our hospital matcher maintains real-time registry of nearby facilities. It scores hospitals '
    'using 5 factors: travel time (30%), bed availability (25%), specialist match (20%), '
    'urgency capability (15%), and facility equipment (10%). A cardiac patient goes to a hospital '
    'with available ICU and cardiologist on duty — not just the nearest one."'
)

doc.add_heading("🎬 Scene 7: API Demonstration (45 seconds)", level=2)
doc.add_paragraph("ACTION:", style='List Bullet').runs[0].bold = True
doc.add_paragraph("  • Show API endpoints in browser:")
doc.add_paragraph("    - http://localhost:3000/api/health")
doc.add_paragraph("    - http://localhost:3000/api/incidents")
doc.add_paragraph("    - http://localhost:3000/api/ambulances")

doc.add_paragraph("NARRATION:", style='List Bullet').runs[0].bold = True
doc.add_paragraph(
    '"Our RESTful API enables integration with existing EMS systems — 108, 102, and private fleets. '
    'The WebSocket server pushes real-time updates to all connected clients including mobile apps '
    'for drivers and citizens."'
)

doc.add_heading("🎬 Scene 8: Conclusion (30 seconds)", level=2)
doc.add_paragraph("ACTION:", style='List Bullet').runs[0].bold = True
doc.add_paragraph("  • Return to dashboard, show the complete system")

doc.add_paragraph("NARRATION:", style='List Bullet').runs[0].bold = True
doc.add_paragraph(
    '"Our AI-Powered Ambulance Dispatch System can reduce emergency response times by up to 40% '
    'through intelligent routing and automated signal preemption. By ensuring the right patient '
    'reaches the right hospital in time, we can save thousands of lives during the golden hour. '
    'Thank you for watching!"'
)

doc.add_page_break()

# ============= SECTION 5: FEATURES =============
doc.add_heading("5. FEATURE HIGHLIGHTS", level=1)

features = [
    ("🚑 Real-Time Optimal Dispatch", 
     "AI selects best ambulance based on live GPS, traffic, vehicle type, and crew availability"),
    ("🚦 Green Corridor Activation",
     "Automatic traffic signal preemption using siren detection (audio FFT + camera OCR)"),
    ("🏥 Intelligent Hospital Matching",
     "5-factor scoring algorithm matches patient condition to appropriate facility"),
    ("📍 Live GPS Tracking",
     "Real-time ambulance positions updated every 5 seconds"),
    ("📊 Unified Dashboard",
     "Single command center for government and private ambulance fleets"),
    ("📱 Citizen Mobile App",
     "One-tap emergency with automatic GPS, live ambulance tracking"),
    ("🔄 Dynamic Route Recalculation",
     "Routes update every 30 seconds based on traffic conditions"),
    ("📈 Predictive Positioning",
     "ML-based demand forecasting for proactive fleet deployment"),
    ("📡 Offline Mode",
     "SMS/USSD fallback and cached maps for rural/low-connectivity areas"),
    ("🔔 Hospital Pre-Alert",
     "Automatic notification to ER with patient vitals and ETA"),
]

for title, desc in features:
    doc.add_heading(title, level=2)
    doc.add_paragraph(desc)

doc.add_page_break()

# ============= SECTION 6: API REFERENCE =============
doc.add_heading("6. API ENDPOINTS REFERENCE", level=1)

doc.add_heading("Core Endpoints", level=2)
table = doc.add_table(rows=8, cols=3)
table.style = 'Table Grid'

headers = ["Method", "Endpoint", "Description"]
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
    table.rows[0].cells[i].paragraphs[0].runs[0].bold = True

api_data = [
    ("GET", "/api/health", "Health check"),
    ("POST", "/api/auth/login", "User authentication"),
    ("GET", "/api/incidents", "List all incidents"),
    ("GET", "/api/ambulances", "List all ambulances"),
    ("GET", "/api/hospitals", "List hospitals with capacity"),
    ("POST", "/api/assignments", "Create ambulance assignment"),
    ("POST", "/api/incidents", "Create new incident"),
]
for i, (method, endpoint, desc) in enumerate(api_data, 1):
    table.rows[i].cells[0].text = method
    table.rows[i].cells[1].text = endpoint
    table.rows[i].cells[2].text = desc

doc.add_heading("Green Corridor Endpoints", level=2)
table2 = doc.add_table(rows=5, cols=3)
table2.style = 'Table Grid'

for i, h in enumerate(headers):
    table2.rows[0].cells[i].text = h
    table2.rows[0].cells[i].paragraphs[0].runs[0].bold = True

gc_data = [
    ("GET", "/api/siren-detection/status", "Detection system status"),
    ("GET", "/api/siren-detection/traffic", "Current traffic light state"),
    ("POST", "/api/green-corridor/activate", "Activate green corridor"),
    ("POST", "/api/green-corridor/deactivate", "Deactivate corridor"),
]
for i, (method, endpoint, desc) in enumerate(gc_data, 1):
    table2.rows[i].cells[0].text = method
    table2.rows[i].cells[1].text = endpoint
    table2.rows[i].cells[2].text = desc

doc.add_page_break()

# ============= SECTION 7: TECH STACK =============
doc.add_heading("7. TECHNOLOGY STACK", level=1)

doc.add_heading("Backend", level=2)
backend_tech = [
    "Node.js + Express.js - API Server",
    "Python + FastAPI - Siren Detection Service",
    "PostgreSQL + PostGIS - Geospatial Database",
    "Redis - Caching & Pub/Sub",
    "Socket.io - Real-time WebSocket",
    "OSRM - Open Source Routing Machine",
]
for t in backend_tech:
    doc.add_paragraph(t, style='List Bullet')

doc.add_heading("Frontend", level=2)
frontend_tech = [
    "React 18 + Vite - Dashboard SPA",
    "Material-UI (MUI) - Component Library",
    "Leaflet.js - Interactive Maps",
    "Chart.js - Analytics Visualization",
]
for t in frontend_tech:
    doc.add_paragraph(t, style='List Bullet')

doc.add_heading("Mobile", level=2)
mobile_tech = [
    "React Native + Expo - Cross-platform Apps",
    "React Native Maps - GPS & Navigation",
]
for t in mobile_tech:
    doc.add_paragraph(t, style='List Bullet')

doc.add_heading("AI/ML", level=2)
ai_tech = [
    "NumPy + SciPy - Audio FFT Analysis",
    "OpenCV + Tesseract - Camera OCR Detection",
    "Scikit-learn - Demand Prediction",
]
for t in ai_tech:
    doc.add_paragraph(t, style='List Bullet')

doc.add_page_break()

# ============= SECTION 8: FUTURE =============
doc.add_heading("8. FUTURE ENHANCEMENTS", level=1)

future = [
    ("IoT Integration", "Connect with ambulance vitals monitors for automatic patient data transmission"),
    ("Voice AI Dispatch", "Natural language processing for emergency call intake"),
    ("Drone Coordination", "Medical drone deployment for rural/inaccessible areas"),
    ("Blockchain Audit", "Immutable emergency response records for accountability"),
    ("Multi-City Scaling", "Federated architecture for pan-India deployment"),
]
for title, desc in future:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(f"{title}: ").bold = True
    p.add_run(desc)

doc.add_page_break()

# ============= SECTION 9: TEAM =============
doc.add_heading("9. TEAM INFORMATION", level=1)
doc.add_paragraph("[Add your team details here]")
doc.add_paragraph()

table3 = doc.add_table(rows=5, cols=3)
table3.style = 'Table Grid'
team_headers = ["Name", "Role", "Contact"]
for i, h in enumerate(team_headers):
    table3.rows[0].cells[i].text = h
    table3.rows[0].cells[i].paragraphs[0].runs[0].bold = True

for i in range(1, 5):
    table3.rows[i].cells[0].text = f"Team Member {i}"
    table3.rows[i].cells[1].text = "Developer"
    table3.rows[i].cells[2].text = "email@example.com"

doc.add_paragraph()
doc.add_paragraph()

# Final note
final = doc.add_paragraph()
final.alignment = WD_ALIGN_PARAGRAPH.CENTER
final.add_run("━" * 50)
doc.add_paragraph()
closing = doc.add_paragraph()
closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
closing.add_run("🚑 SAVING LIVES, ONE DISPATCH AT A TIME 🚑").bold = True

# Save document
output_path = r"C:\Users\Admin\EVERYTHING-AMBULANCE-FEATURES\AMBULANCE_DISPATCH_DEMO_GUIDE.docx"
doc.save(output_path)
print(f"✅ Document created: {output_path}")
